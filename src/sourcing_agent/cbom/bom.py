"""BOM and CBOM files: read a BOM (CSV or .xlsx), write a CBOM, read CBOMs back."""

# BOM and CBOM files.
import csv
# Corrupt .xlsx errors.
import zipfile
# File paths.
from pathlib import Path

# Excel BOM files.
import openpyxl
# Unsupported workbook errors.
from openpyxl.utils.exceptions import InvalidFileException

# Error for bad input files.
from ..core.errors import InputError

# BOM input columns, in order.
BOM_COLUMNS = ["part_id", "description", "category", "quantity", "spec", "mfr_part_number", "notes"]
# BOM columns that must have a value.
REQUIRED_COLUMNS = ["part_id", "description", "category", "quantity"]
# CBOM output columns, in order.
CBOM_COLUMNS = BOM_COLUMNS + [
    "status", "vendor", "vendor_part_number", "url", "pack_size", "order_packs", "order_qty",
    "pack_price", "extended_price", "effective_unit_price", "currency", "quotes_compared",
    "quoted_at", "sourcing_notes",
]
# BOM file extensions read as Excel workbooks.
XLSX_SUFFIXES = {".xlsx", ".xlsm"}
# Workbook sheet holding the BOM.
XLSX_SHEET = "BOM"


def cell_text(value: object) -> str:
    """Return an Excel cell value as BOM text."""
    # Empty cell.
    if value is None:
        return ""
    # Whole-number floats print without ".0".
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    # Anything else as plain text.
    return str(value)


def read_csv_rows(path: Path) -> tuple[list[str], list[tuple[int, dict]]]:
    """Return the header and (line, row) pairs from a BOM CSV."""
    # Read all rows.
    try:
        with path.open(newline="", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            rows = list(reader)
            header = reader.fieldnames or []
    except OSError as e:
        raise InputError(f"cannot read {path}: {e}") from e
    # Line 2 is the first data row.
    return header, list(enumerate(rows, start=2))


def read_xlsx_rows(path: Path) -> tuple[list[str], list[tuple[int, dict]]]:
    """Return the header and (row number, row) pairs from the BOM sheet of a workbook."""
    # Open for reading, with formula results instead of formulas.
    try:
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except (OSError, zipfile.BadZipFile, InvalidFileException) as e:
        raise InputError(f"cannot read {path}: {e}") from e
    # The BOM sheet, or the first sheet if there is none.
    sheet = book[XLSX_SHEET] if XLSX_SHEET in book.sheetnames else book.worksheets[0]
    # All cells as trimmed text.
    values = [[cell_text(v).strip() for v in r] for r in sheet.iter_rows(values_only=True)]
    # Release the file.
    book.close()
    # Row 1 is the header.
    header = values[0] if values else []
    # Data rows by sheet row number, skipping blank rows.
    return header, [(n, dict(zip(header, v))) for n, v in enumerate(values[1:], start=2) if any(v)]


def load_bom(path: Path, categories: list[str]) -> list[dict]:
    """Read a BOM CSV or .xlsx, check every row, and return rows with all BOM columns."""
    # Read the header and numbered rows for the file type.
    header, numbered = read_xlsx_rows(path) if path.suffix.lower() in XLSX_SUFFIXES else read_csv_rows(path)
    # Required columns must be in the header.
    missing = [c for c in REQUIRED_COLUMNS if c not in header]
    if missing:
        raise InputError(f"{path}: missing columns {missing}")
    # Part IDs seen so far.
    seen = set()
    # Check each row.
    for line, row in numbered:
        # Fill absent optional columns and trim values.
        row.update({c: (row.get(c) or "").strip() for c in BOM_COLUMNS})
        # Required values must be present.
        if not all(row[c] for c in REQUIRED_COLUMNS):
            raise InputError(f"{path}:{line}: empty required value")
        # Part IDs must be unique.
        if row["part_id"] in seen:
            raise InputError(f"{path}:{line}: duplicate part_id {row['part_id']}")
        # Remember this part ID.
        seen.add(row["part_id"])
        # Category must be in suppliers.toml.
        if row["category"] not in categories:
            raise InputError(f"{path}:{line}: unknown category '{row['category']}'")
        # Quantity must be a positive integer.
        if not row["quantity"].isdigit() or int(row["quantity"]) < 1:
            raise InputError(f"{path}:{line}: quantity must be a positive integer")
    # Rows with only the BOM columns.
    return [{c: row[c] for c in BOM_COLUMNS} for _, row in numbered]


def write_cbom(path: Path, rows: list[dict]) -> None:
    """Write the CBOM rows to a CSV file."""
    # Create the output folder if needed.
    path.parent.mkdir(parents=True, exist_ok=True)
    # Write header and rows.
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=CBOM_COLUMNS)
        writer.writeheader()
        writer.writerows(rows)


def read_cboms(paths: list[Path]) -> list[dict]:
    """Read the rows of every CBOM file."""
    # All rows, in file order.
    rows = []
    for path in paths:
        with path.open(newline="", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            # A CBOM needs the columns later steps rely on.
            missing = [c for c in ("category", "status", "vendor") if c not in (reader.fieldnames or [])]
            if missing:
                raise InputError(f"{path}: not a CBOM (missing columns {missing})")
            rows.extend(reader)
    # Combined rows.
    return rows
