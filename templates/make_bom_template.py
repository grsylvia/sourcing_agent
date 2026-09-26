"""Build templates/bom_template.xlsx from the BOM columns and suppliers.toml categories."""

# Supplier list file.
import tomllib
# File paths.
from pathlib import Path

# Workbook writer.
from openpyxl import Workbook
# Header cell notes.
from openpyxl.comments import Comment
# Cell styles.
from openpyxl.styles import Alignment, Font, PatternFill
# Column number to letter.
from openpyxl.utils import get_column_letter
# Dropdown and number checks.
from openpyxl.worksheet.datavalidation import DataValidation

# BOM columns, required columns, and sheet name the loader reads.
from sourcing_agent.orchestrator import BOM_COLUMNS, REQUIRED_COLUMNS, XLSX_SHEET

# Project root folder.
ROOT = Path(__file__).resolve().parents[1]
# Template output path.
OUT = ROOT / "templates" / "bom_template.xlsx"
# Blank input rows formatted on the BOM sheet.
INPUT_ROWS = 500
# Meaning and example for each BOM column.
COLUMN_INFO = {
    "part_id": ("Unique ID for the line", "F-001"),
    "description": ("What the part is", "M3x8 socket head cap screw"),
    "category": ("Supplier category (dropdown)", "fasteners"),
    "quantity": ("Units needed, whole number > 0", 40),
    "spec": ("Material, size, standard, rating", "ISO 4762, A2 stainless"),
    "mfr_part_number": ("Manufacturer part number, if known", ""),
    "notes": ("Anything else the sourcer should know", ""),
}
# Column widths by column name.
WIDTHS = {"part_id": 12, "description": 36, "category": 14, "quantity": 10, "spec": 36, "mfr_part_number": 20, "notes": 36}
# Columns kept as text so IDs like 00123 keep leading zeros.
TEXT_COLUMNS = {"part_id", "spec", "mfr_part_number", "notes", "description"}

# Base font for every cell.
FONT = Font(name="Arial", size=10)
# Bold white header font.
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
# Title font on the instructions sheet.
TITLE_FONT = Font(name="Arial", size=14, bold=True)
# Section heading font on the instructions sheet.
BOLD = Font(name="Arial", size=10, bold=True)
# Dark blue fill for required headers.
REQUIRED_FILL = PatternFill("solid", fgColor="1F4E78")
# Grey fill for optional headers.
OPTIONAL_FILL = PatternFill("solid", fgColor="808080")
# Light yellow fill for cells to fill in.
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")


def load_categories() -> list[str]:
    """Return the allowed categories from suppliers.toml."""
    # Parse the supplier list.
    with (ROOT / "suppliers.toml").open("rb") as f:
        return tomllib.load(f)["categories"]


def input_range(name: str) -> str:
    """Return the input cell range of a BOM column, e.g. C2:C501."""
    # Column letter for this BOM column.
    letter = get_column_letter(BOM_COLUMNS.index(name) + 1)
    # Rows below the header.
    return f"{letter}2:{letter}{INPUT_ROWS + 1}"


def build_bom_sheet(sheet, categories: list[str]) -> None:
    """Write the header, input area, and checks on the BOM sheet."""
    # Header row with required columns dark and optional columns grey.
    for col, name in enumerate(BOM_COLUMNS, start=1):
        cell = sheet.cell(row=1, column=col, value=name)
        cell.font = HEADER_FONT
        cell.fill = REQUIRED_FILL if name in REQUIRED_COLUMNS else OPTIONAL_FILL
        cell.comment = Comment(f"{COLUMN_INFO[name][0]}. {'Required' if name in REQUIRED_COLUMNS else 'Optional'}.", "BOM template")
        sheet.column_dimensions[cell.column_letter].width = WIDTHS[name]
    # Yellow input cells with the base font.
    for row in sheet.iter_rows(min_row=2, max_row=INPUT_ROWS + 1, max_col=len(BOM_COLUMNS)):
        for cell in row:
            cell.font = FONT
            cell.fill = INPUT_FILL
            cell.number_format = "@" if BOM_COLUMNS[cell.column - 1] in TEXT_COLUMNS else "General"
    # Keep the header visible while scrolling.
    sheet.freeze_panes = "A2"
    # Category dropdown limited to suppliers.toml categories.
    category = DataValidation(type="list", formula1=f'"{",".join(categories)}"', allow_blank=True, showErrorMessage=True, errorTitle="Unknown category", error="Pick a category from the list.")
    category.add(input_range("category"))
    sheet.add_data_validation(category)
    # Quantity limited to whole numbers of at least 1.
    quantity = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="1", allow_blank=True, showErrorMessage=True, errorTitle="Bad quantity", error="Enter a whole number of 1 or more.")
    quantity.add(input_range("quantity"))
    sheet.add_data_validation(quantity)


def build_instructions_sheet(sheet, categories: list[str]) -> None:
    """Write the legend, column guide, example row, and categories."""
    # Lines to write, top to bottom; None leaves a blank row.
    lines = [
        ("BOM template", TITLE_FONT),
        ("Fill in the BOM sheet, save as .xlsx, then run: sourcing run <file>.xlsx", FONT),
        None,
        ("Legend", BOLD),
        ("Yellow cells on the BOM sheet: fill these in, one part per row", FONT),
        ("Dark blue header: required column", FONT),
        ("Grey header: optional column", FONT),
        ("Only the BOM sheet is read; blank rows are skipped; do not rename headers", FONT),
        None,
        ("Columns", BOLD),
    ]
    # Current row.
    r = 1
    # Title, legend, and section heading.
    for line in lines:
        if line:
            sheet.cell(row=r, column=1, value=line[0]).font = line[1]
        r += 1
    # Column guide header.
    for col, label in enumerate(["Column", "Required", "Meaning", "Example"], start=1):
        sheet.cell(row=r, column=col, value=label).font = BOLD
    r += 1
    # One guide row per BOM column.
    for name in BOM_COLUMNS:
        meaning, example = COLUMN_INFO[name]
        for col, value in enumerate([name, "Yes" if name in REQUIRED_COLUMNS else "No", meaning, example], start=1):
            sheet.cell(row=r, column=col, value=value).font = FONT
        r += 1
    # Example row heading.
    r += 1
    sheet.cell(row=r, column=1, value="Example BOM row (reference only; not read)").font = BOLD
    r += 1
    # Example header and values laid out like the BOM sheet.
    for col, name in enumerate(BOM_COLUMNS, start=1):
        sheet.cell(row=r, column=col, value=name).font = BOLD
        sheet.cell(row=r + 1, column=col, value=COLUMN_INFO[name][1]).font = FONT
    r += 3
    # Allowed categories.
    sheet.cell(row=r, column=1, value="Categories (from suppliers.toml)").font = BOLD
    r += 1
    for name in categories:
        sheet.cell(row=r, column=1, value=name).font = FONT
        r += 1
    # Column widths for readability.
    for letter, width in zip("ABCDEFG", [18, 36, 36, 28, 36, 20, 36]):
        sheet.column_dimensions[letter].width = width
    # Align text to the top.
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top")


def main() -> None:
    """Build and save the template."""
    # Allowed categories.
    categories = load_categories()
    # New workbook; the first sheet is the BOM.
    book = Workbook()
    # BOM sheet the loader reads.
    bom = book.active
    bom.title = XLSX_SHEET
    build_bom_sheet(bom, categories)
    # Instructions sheet.
    build_instructions_sheet(book.create_sheet("Instructions"), categories)
    # Save the template.
    book.save(OUT)
    print(f"Wrote {OUT}")


# Allow python templates/make_bom_template.py.
if __name__ == "__main__":
    main()
